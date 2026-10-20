from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.units import mm
from reportlab.platypus import Spacer, Table

from quotes.pdf import build_document, fmt, p, qty, styles, table_style


def price_list_xlsx(price_list, items) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = price_list.code[:31]
    ws.append([f"Lista {price_list.code} · {price_list.name}"])
    ws.append(
        [
            f"Moneda: {price_list.currency}",
            f"Vigencia: {price_list.valid_from:%d/%m/%Y} – "
            + (f"{price_list.valid_until:%d/%m/%Y}" if price_list.valid_until else "sin fecha final"),
        ]
    )
    ws.append([])
    ws.append(["SKU", "Producto", "Unidad", "Cantidad mínima", "Precio unitario", "Moneda"])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A2540")
    ws["A1"].font = Font(bold=True, size=13, color="0A2540")
    for item in items:
        ws.append(
            [
                item.product.sku,
                item.product.name,
                item.product.unit.code,
                item.min_quantity,
                item.unit_price,
                price_list.currency,
            ]
        )
    for col, width in zip("ABCDEF", [16, 44, 8, 16, 16, 8], strict=True):
        ws.column_dimensions[col].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def price_list_pdf(price_list, items) -> bytes:
    st = styles()
    validity = f"{price_list.valid_from:%d/%m/%Y} – " + (
        f"{price_list.valid_until:%d/%m/%Y}" if price_list.valid_until else "sin fecha final"
    )
    story = [
        p(f"{price_list.name}", st["h"]),
        p(
            f"Clave {price_list.code} · Moneda {price_list.currency} · Vigencia {validity}"
            + (f" · Cliente: {price_list.customer.name}" if price_list.customer else ""),
            st["small"],
        ),
        Spacer(1, 5 * mm),
    ]
    rows = [["SKU", "Producto", "Unidad", "Cant. mínima", f"Precio ({price_list.currency})"]]
    for item in items:
        rows.append(
            [
                item.product.sku,
                p(item.product.name, st["normal"]),
                item.product.unit.code,
                qty(item.min_quantity),
                fmt(item.unit_price, 4 if item.unit_price < 10 else 2),
            ]
        )
    table = Table(rows, colWidths=[30 * mm, 80 * mm, 18 * mm, 25 * mm, 32 * mm], repeatRows=1)
    style = table_style()
    style.add("ALIGN", (3, 0), (-1, -1), "RIGHT")
    table.setStyle(style)
    story += [
        table,
        Spacer(1, 5 * mm),
        p("Precios sin IVA. Sujetos a cambio sin previo aviso fuera de la vigencia indicada.", st["small"]),
    ]
    return build_document(story, f"Lista de precios {price_list.code}")
