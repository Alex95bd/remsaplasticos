"""Importación y exportación del catálogo de productos en Excel (.xlsx)."""

import datetime
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from audit.models import AuditLog
from audit.services import diff_snapshots, log_action, model_snapshot
from core.choices import Currency

from .models import Category, Product, UnitOfMeasure, normalize_sku

# (columna, campo del modelo, obligatoria, descripción)
COLUMNS = [
    ("sku", "sku", True, "Código único del producto. Se convierte a mayúsculas."),
    ("nombre", "name", True, "Nombre comercial."),
    ("descripcion", "description", False, "Descripción libre."),
    ("categoria", "category", True, "Nombre exacto de una categoría existente."),
    ("tipo", "product_type", False, "Tipo de producto. Vacío = Producto terminado."),
    ("unidad", "unit", True, "Clave o nombre de una unidad de medida existente (ej. PZA)."),
    ("presentacion", "presentation", False, "Ej. Caja con 50 piezas."),
    ("capacidad_ml", "capacity_ml", False, "Número en mililitros."),
    ("alto_mm", "height_mm", False, "Número en milímetros."),
    ("diametro_mm", "diameter_mm", False, "Número en milímetros."),
    ("peso_g", "weight_g", False, "Número en gramos."),
    ("precio_referencia", "reference_price", False, "Precio unitario sin impuestos, hasta 4 decimales."),
    ("moneda", "currency", False, "MXN o USD. Vacío = MXN."),
    ("codigo_externo", "external_code", False, "Clave del artículo en Intelisis (opcional, única)."),
    ("activo", "is_active", False, "SI o NO. Vacío = SI."),
]
COLUMN_KEYS = [column[0] for column in COLUMNS]
REQUIRED_COLUMNS = [column[0] for column in COLUMNS if column[2]]
PRICE_COLUMNS = {"precio_referencia", "moneda"}
DECIMAL_COLUMNS = {"capacidad_ml": 2, "alto_mm": 2, "diametro_mm": 2, "peso_g": 2, "precio_referencia": 4}
HEADER_ALIASES = {
    "codigo": "sku",
    "clave": "sku",
    "nombre_comercial": "nombre",
    "tipo_de_producto": "tipo",
    "unidad_de_medida": "unidad",
    "precio": "precio_referencia",
    "precio_de_referencia": "precio_referencia",
    "codigo_intelisis": "codigo_externo",
}
TRUE_VALUES = {"si", "s", "1", "true", "verdadero", "activo", "x"}
FALSE_VALUES = {"no", "n", "0", "false", "falso", "inactivo"}


class ImportFileError(Exception):
    """Error que impide procesar el archivo completo."""


def _plain(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


def normalize_header(value) -> str:
    key = _plain(value)
    return HEADER_ALIASES.get(key, key)


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "SI" if value else "NO"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(Decimal(str(value)))
    if isinstance(value, datetime.datetime | datetime.date):
        return value.isoformat()
    return str(value).strip()


def read_workbook(fileobj, max_rows: int):
    """Lee la primera hoja y devuelve ``(filas, columnas_ignoradas)``."""
    try:
        workbook = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza distintos tipos de error para archivos inválidos
        raise ImportFileError("No se pudo leer el archivo. Verifica que sea un Excel .xlsx válido.") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not any(header):
            raise ImportFileError("La primera fila debe contener los nombres de las columnas.")
        keys = [normalize_header(value) for value in header]
        repeated = sorted({key for key in keys if key and keys.count(key) > 1})
        if repeated:
            raise ImportFileError(f"Columnas repetidas: {', '.join(repeated)}.")
        missing = [key for key in REQUIRED_COLUMNS if key not in keys]
        if missing:
            raise ImportFileError(f"Faltan columnas obligatorias: {', '.join(missing)}.")
        ignored = [str(value) for value, key in zip(header, keys, strict=False) if value and key not in COLUMN_KEYS]
        data = []
        for index, values in enumerate(rows, start=2):
            record = {key: _cell_to_text(value) for key, value in zip(keys, values, strict=False) if key in COLUMN_KEYS}
            if not any(record.values()):
                continue
            if len(data) >= max_rows:
                raise ImportFileError(f"El archivo excede el máximo de {max_rows} filas por importación.")
            data.append({"row": index, "values": record})
        if not data:
            raise ImportFileError("El archivo no contiene filas con datos.")
        return data, ignored
    finally:
        workbook.close()


def _type_lookup():
    lookup = {}
    for value, label in Product.ProductType.choices:
        lookup[_plain(value)] = value
        lookup[_plain(label)] = value
        lookup[_plain(label.split("(")[0])] = value
    lookup["preforma"] = Product.ProductType.SEMI_FINISHED
    return lookup


def _parse_decimal(text: str, places: int):
    cleaned = text.replace("$", "").replace(",", "").replace(" ", "")
    if not cleaned:
        return None, None
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return None, f"«{text}» no es un número válido."
    if not value.is_finite():
        return None, f"«{text}» no es un número válido."
    if value < 0:
        return None, "No se permiten valores negativos."
    return value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP), None


def _parse_bool(text: str):
    plain = _plain(text)
    if not plain:
        return True, None
    if plain in TRUE_VALUES:
        return True, None
    if plain in FALSE_VALUES:
        return False, None
    return None, f"«{text}» no es válido; usa SI o NO."


@dataclass
class RowResult:
    row: int
    values: dict
    action: str = "crear"  # crear | actualizar | sin_cambios | error
    errors: list = field(default_factory=list)
    changes: dict = field(default_factory=dict)
    data: dict = field(default_factory=dict)
    product_id: int | None = None

    @property
    def is_valid(self):
        return not self.errors


@dataclass
class ImportPreview:
    rows: list

    def count(self, action):
        return sum(1 for row in self.rows if row.action == action)

    @property
    def to_create(self):
        return self.count("crear")

    @property
    def to_update(self):
        return self.count("actualizar")

    @property
    def unchanged(self):
        return self.count("sin_cambios")

    @property
    def errors(self):
        return self.count("error")

    @property
    def applicable(self):
        return self.to_create + self.to_update


def validate_rows(rows, *, update_existing: bool, include_prices: bool = True) -> ImportPreview:
    categories = {category.name.lower(): category for category in Category.objects.all()}
    units = {}
    for unit in UnitOfMeasure.objects.all():
        units[unit.code.lower()] = unit
        units.setdefault(unit.name.lower(), unit)
    types = _type_lookup()
    skus = [normalize_sku(row["values"].get("sku")) for row in rows]
    existing = {product.sku: product.pk for product in Product.objects.filter(sku__in=skus).only("pk", "sku")}
    audit_fields = [column[1] for column in COLUMNS if include_prices or column[0] not in PRICE_COLUMNS]
    seen_skus, seen_external = set(), set()
    results = []

    for row in rows:
        values = row["values"]
        result = RowResult(row=row["row"], values=values)
        errors = result.errors
        data = {}

        sku = normalize_sku(values.get("sku"))
        if not sku:
            errors.append("El SKU es obligatorio.")
        elif sku in seen_skus:
            errors.append("SKU repetido dentro del archivo.")
        seen_skus.add(sku)
        data["sku"] = sku

        data["name"] = " ".join(values.get("nombre", "").split())
        if not data["name"]:
            errors.append("El nombre es obligatorio.")
        data["description"] = values.get("descripcion", "")
        data["presentation"] = values.get("presentacion", "")

        category_name = values.get("categoria", "").strip()
        category = categories.get(category_name.lower())
        if not category:
            errors.append(f"La categoría «{category_name}» no existe; créala antes de importar.")
        elif not category.is_active:
            errors.append(f"La categoría «{category.name}» está inactiva.")
        data["category"] = category

        unit_name = values.get("unidad", "").strip()
        unit = units.get(unit_name.lower())
        if not unit:
            errors.append(f"La unidad «{unit_name}» no existe.")
        elif not unit.is_active:
            errors.append(f"La unidad «{unit.code}» está inactiva.")
        data["unit"] = unit

        type_text = values.get("tipo", "")
        data["product_type"] = types.get(_plain(type_text)) if type_text else Product.ProductType.FINISHED
        if not data["product_type"]:
            errors.append(f"Tipo de producto «{type_text}» no reconocido.")

        for column, places in DECIMAL_COLUMNS.items():
            if column in PRICE_COLUMNS and not include_prices:
                continue
            number, error = _parse_decimal(values.get(column, ""), places)
            if error:
                errors.append(f"{column}: {error}")
            data[dict((c[0], c[1]) for c in COLUMNS)[column]] = number

        if include_prices:
            currency = values.get("moneda", "").strip().upper() or Currency.MXN
            if currency not in Currency.values:
                errors.append(f"Moneda «{currency}» no válida; usa MXN o USD.")
            data["currency"] = currency

        external = values.get("codigo_externo", "").strip()
        if external and external in seen_external:
            errors.append("Código externo repetido dentro del archivo.")
        if external:
            seen_external.add(external)
        data["external_code"] = external

        active, error = _parse_bool(values.get("activo", ""))
        if error:
            errors.append(f"activo: {error}")
        data["is_active"] = active

        product_id = existing.get(sku)
        if product_id and not update_existing:
            errors.append(
                "Ya existe un producto con este SKU. Marca «Actualizar productos existentes» para modificarlo."
            )

        if not errors:
            instance = Product.objects.get(pk=product_id) if product_id else Product()
            before = model_snapshot(instance, audit_fields) if product_id else {}
            for name, value in data.items():
                setattr(instance, name, value)
            try:
                instance.full_clean()
            except ValidationError as exc:
                for name, messages in exc.message_dict.items():
                    label = Product._meta.get_field(name).verbose_name if name != "__all__" else "Producto"
                    errors.extend(f"{label}: {message}" for message in messages)
            if product_id and not errors:
                result.changes = diff_snapshots(before, model_snapshot(instance, audit_fields))

        result.data = data
        result.product_id = product_id
        if errors:
            result.action = "error"
        elif product_id:
            result.action = "actualizar" if result.changes else "sin_cambios"
        results.append(result)
    return ImportPreview(rows=results)


def apply_import(rows, *, update_existing: bool, include_prices: bool, request, filename: str = ""):
    """Aplica las filas válidas en una sola transacción y devuelve ``(preview, creados, actualizados)``."""
    user = request.user
    try:
        with transaction.atomic():
            preview = validate_rows(rows, update_existing=update_existing, include_prices=include_prices)
            created, updated = [], []
            for result in preview.rows:
                if result.action == "crear":
                    product = Product(**result.data, created_by=user, updated_by=user)
                    product.save()
                    created.append(product.sku)
                elif result.action == "actualizar":
                    product = Product.objects.select_for_update().get(pk=result.product_id)
                    for name, value in result.data.items():
                        setattr(product, name, value)
                    product.updated_by = user
                    product.save()
                    updated.append(product.sku)
            log_action(
                AuditLog.Action.IMPORT,
                request=request,
                message=(
                    f"Importación de productos «{filename[:80]}»: {len(created)} altas, {len(updated)} "
                    f"actualizaciones, {preview.errors} filas omitidas por errores."
                ),
                changes={"creados": created[:500], "actualizados": updated[:500]},
            )
    except IntegrityError as exc:
        raise ImportFileError(
            "Otro usuario modificó el catálogo al mismo tiempo. Vuelve a cargar el archivo para revisar los cambios."
        ) from exc
    return preview, created, updated


HEADER_FILL = PatternFill("solid", fgColor="0A2540")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_header(sheet, headers):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(14, len(header) + 4)


def _workbook_bytes(workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _columns(include_prices):
    return [column for column in COLUMNS if include_prices or column[0] not in PRICE_COLUMNS]


def export_products(queryset, *, include_prices: bool) -> bytes:
    columns = _columns(include_prices)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Productos"
    _write_header(sheet, [column[0] for column in columns])
    type_labels = dict(Product.ProductType.choices)
    for product in queryset.select_related("category", "unit"):
        row = []
        for key, attr, _required, _help in columns:
            value = getattr(product, attr)
            if key == "categoria":
                value = product.category.name
            elif key == "unidad":
                value = product.unit.code
            elif key == "tipo":
                value = type_labels.get(value, value)
            elif key == "activo":
                value = "SI" if value else "NO"
            row.append(value)
        sheet.append(row)
    return _workbook_bytes(workbook)


def import_template(*, include_prices: bool) -> bytes:
    columns = _columns(include_prices)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Productos"
    _write_header(sheet, [column[0] for column in columns])

    help_sheet = workbook.create_sheet("Instrucciones")
    _write_header(help_sheet, ["columna", "obligatoria", "descripción"])
    for key, _attr, required, description in columns:
        help_sheet.append([key, "SI" if required else "NO", description])
    help_sheet.column_dimensions["C"].width = 70
    help_sheet.append([])
    help_sheet.append(["Tipos válidos", "", ", ".join(label for _value, label in Product.ProductType.choices)])
    help_sheet.append(
        ["Unidades", "", ", ".join(UnitOfMeasure.objects.filter(is_active=True).values_list("code", flat=True))]
    )
    help_sheet.append(
        ["Categorías", "", ", ".join(Category.objects.filter(is_active=True).values_list("name", flat=True))]
    )
    return _workbook_bytes(workbook)
