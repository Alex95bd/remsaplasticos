from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.roles import ROLE_ADMIN, ROLE_READONLY, ROLE_SALES, ROLE_WAREHOUSE
from audit.models import AuditLog
from catalog.models import Category, Product, UnitOfMeasure
from core.tests.factories import make_category, make_product, make_user


def xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("productos.xlsx", buffer.getvalue())


def product_payload(**overrides):
    data = {
        "sku": " gar-20l ",
        "name": "Garrafa PEAD 20 L",
        "category": Category.objects.first().pk,
        "product_type": "terminado",
        "unit": UnitOfMeasure.objects.get(code="PZA").pk,
        "reference_price": "35.50",
        "currency": "MXN",
        "is_active": "on",
    }
    data.update(overrides)
    return data


class ProductModelTests(TestCase):
    def test_base_units_exist(self):
        self.assertTrue(UnitOfMeasure.objects.filter(code="PZA").exists())

    def test_sku_is_normalized(self):
        product = make_product(" abc-1 ")
        self.assertEqual(product.sku, "ABC-1")

    def test_category_name_unique_case_insensitive(self):
        make_category("Tapas")
        from django.core.exceptions import ValidationError

        with self.assertRaises(ValidationError):
            Category(name="tapas").full_clean()


class ProductViewTests(TestCase):
    def setUp(self):
        self.category = make_category()
        self.admin = make_user("admin", ROLE_ADMIN)

    def test_create_product_and_audit(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("catalog:product_create"), product_payload())
        product = Product.objects.get(sku="GAR-20L")
        self.assertRedirects(response, reverse("catalog:product_detail", args=[product.pk]))
        self.assertEqual(product.reference_price, Decimal("35.50"))
        self.assertEqual(product.created_by, self.admin)
        entry = AuditLog.objects.get(action="create", object_id=str(product.pk))
        self.assertEqual(entry.changes["reference_price"]["despues"], "35.5000")

    def test_duplicate_sku_rejected(self):
        make_product("GAR-20L")
        self.client.force_login(self.admin)
        response = self.client.post(reverse("catalog:product_create"), product_payload(sku="gar-20l"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("sku", response.context["form"].errors)

    def test_negative_price_rejected(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("catalog:product_create"), product_payload(reference_price="-1"))
        self.assertIn("reference_price", response.context["form"].errors)

    def test_update_logs_changes(self):
        product = make_product("X-1")
        self.client.force_login(self.admin)
        self.client.post(
            reverse("catalog:product_update", args=[product.pk]),
            product_payload(sku="X-1", name="Nuevo nombre", category=self.category.pk, reference_price="40"),
        )
        entry = AuditLog.objects.get(action="update", object_id=str(product.pk))
        self.assertEqual(entry.changes["name"], {"antes": "Garrafa de prueba 20 L", "despues": "Nuevo nombre"})

    def test_toggle_active_requires_post_and_permission(self):
        product = make_product("X-2")
        self.client.force_login(make_user("vend", ROLE_SALES))
        self.assertEqual(self.client.post(reverse("catalog:product_toggle_active", args=[product.pk])).status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("catalog:product_toggle_active", args=[product.pk])).status_code, 405)
        self.client.post(reverse("catalog:product_toggle_active", args=[product.pk]))
        product.refresh_from_db()
        self.assertFalse(product.is_active)

    def test_readonly_roles_cannot_write(self):
        for role in (ROLE_SALES, ROLE_WAREHOUSE, ROLE_READONLY):
            self.client.force_login(make_user(f"u-{role}", role))
            self.assertEqual(self.client.get(reverse("catalog:product_list")).status_code, 200)
            self.assertEqual(self.client.post(reverse("catalog:product_create"), product_payload()).status_code, 403)
        self.assertFalse(Product.objects.exists())

    def test_prices_hidden_for_warehouse(self):
        make_product("P-1", reference_price=Decimal("123.45"))
        self.client.force_login(make_user("alm", ROLE_WAREHOUSE))
        response = self.client.get(reverse("catalog:product_list"))
        self.assertNotContains(response, "123.45")
        self.client.force_login(make_user("ven", ROLE_SALES))
        self.assertContains(self.client.get(reverse("catalog:product_list")), "123.45")

    def test_search_and_filters(self):
        other = make_category("Tapas")
        make_product("GAR-1", name="Garrafa azul")
        make_product("TAP-1", name="Tapa 28 mm", category=other, product_type="insumo")
        make_product("OLD-1", name="Garrafa vieja", is_active=False)
        self.client.force_login(self.admin)
        url = reverse("catalog:product_list")
        skus = lambda r: [p.sku for p in r.context["products"]]  # noqa: E731
        self.assertEqual(skus(self.client.get(url, {"q": "garrafa"})), ["GAR-1"])
        self.assertEqual(skus(self.client.get(url, {"categoria": other.pk})), ["TAP-1"])
        self.assertEqual(skus(self.client.get(url, {"tipo": "insumo"})), ["TAP-1"])
        self.assertEqual(skus(self.client.get(url, {"estado": "inactivos"})), ["OLD-1"])


class ImportExportTests(TestCase):
    header = ["SKU", "Nombre", "Categoría", "Unidad", "Precio", "Activo"]

    def setUp(self):
        self.category = make_category("Garrafas PEAD")
        self.client.force_login(make_user("admin", ROLE_ADMIN))

    def upload(self, rows, update_existing=False):
        data = {"file": xlsx(rows)}
        if update_existing:
            data["update_existing"] = "on"
        return self.client.post(reverse("catalog:product_import"), data)

    def test_missing_required_columns(self):
        response = self.upload([["sku", "nombre"], ["A", "B"]])
        self.assertContains(response, "Faltan columnas obligatorias: categoria, unidad")

    def test_invalid_file(self):
        response = self.client.post(
            reverse("catalog:product_import"), {"file": SimpleUploadedFile("x.xlsx", b"no es excel")}
        )
        self.assertContains(response, "No se pudo leer el archivo")

    def test_preview_then_confirm(self):
        make_product("EXIST-1", category=self.category)
        response = self.upload(
            [
                self.header,
                ["new-1", "Garrafa 5 L", "garrafas pead", "PZA", 14.5, "si"],
                ["NEW-1", "Duplicada", "Garrafas PEAD", "PZA", 1, "si"],
                ["NEW-2", "Sin categoría", "Inexistente", "PZA", 1, ""],
                ["NEW-3", "Precio malo", "Garrafas PEAD", "PZA", "abc", ""],
                ["EXIST-1", "Existente", "Garrafas PEAD", "PZA", 1, ""],
                [None, None, None, None, None, None],
            ]
        )
        self.assertRedirects(response, reverse("catalog:product_import_preview"))
        preview = self.client.get(reverse("catalog:product_import_preview")).context["preview"]
        self.assertEqual([r.action for r in preview.rows], ["crear", "error", "error", "error", "error"])
        self.assertIn("SKU repetido dentro del archivo.", preview.rows[1].errors)
        self.assertFalse(Product.objects.filter(sku="NEW-1").exists(), "La vista previa no debe guardar")

        response = self.client.post(reverse("catalog:product_import_preview"), {"confirmar": "1"})
        self.assertRedirects(response, reverse("catalog:product_list"))
        product = Product.objects.get(sku="NEW-1")
        self.assertEqual(product.reference_price, Decimal("14.5000"))
        self.assertEqual(Product.objects.count(), 2)
        entry = AuditLog.objects.get(action="import")
        self.assertEqual(entry.changes["creados"], ["NEW-1"])

    def test_update_existing(self):
        make_product("EXIST-1", category=self.category, name="Antes")
        self.upload([self.header, ["EXIST-1", "Después", "Garrafas PEAD", "PZA", 35, "no"]], update_existing=True)
        preview = self.client.get(reverse("catalog:product_import_preview")).context["preview"]
        self.assertEqual(preview.rows[0].action, "actualizar")
        self.assertEqual(preview.rows[0].changes["name"], {"antes": "Antes", "despues": "Después"})
        self.client.post(reverse("catalog:product_import_preview"), {"confirmar": "1"})
        product = Product.objects.get(sku="EXIST-1")
        self.assertEqual(product.name, "Después")
        self.assertFalse(product.is_active)

    def test_cancel_discards(self):
        self.upload([self.header, ["C-1", "Cancelado", "Garrafas PEAD", "PZA", 1, ""]])
        self.client.post(reverse("catalog:product_import_preview"), {"cancelar": "1"})
        self.assertFalse(Product.objects.exists())
        self.assertRedirects(
            self.client.get(reverse("catalog:product_import_preview")), reverse("catalog:product_import")
        )

    def test_import_forbidden_for_sales(self):
        self.client.force_login(make_user("vend", ROLE_SALES))
        self.assertEqual(self.client.get(reverse("catalog:product_import")).status_code, 403)

    def test_export_respects_price_permission(self):
        make_product("E-1", category=self.category, reference_price=Decimal("9.99"))
        response = self.client.get(reverse("catalog:product_export"))
        sheet = load_workbook(BytesIO(response.content)).active
        header = [cell.value for cell in sheet[1]]
        self.assertIn("precio_referencia", header)
        self.assertEqual(sheet.cell(row=2, column=1).value, "E-1")
        self.assertTrue(AuditLog.objects.filter(action="export").exists())
        self.client.force_login(make_user("alm", ROLE_WAREHOUSE))
        self.assertEqual(self.client.get(reverse("catalog:product_export")).status_code, 403)

    def test_template_download(self):
        response = self.client.get(reverse("catalog:product_import_template"))
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Productos", "Instrucciones"])


class ProductApiTests(TestCase):
    def setUp(self):
        self.category = make_category()
        self.product = make_product("API-1", category=self.category, reference_price=Decimal("10"))

    def test_requires_authentication(self):
        self.assertEqual(self.client.get("/api/v1/productos/").status_code, 403)

    def test_warehouse_reads_without_prices_and_cannot_write(self):
        self.client.force_login(make_user("alm", ROLE_WAREHOUSE))
        data = self.client.get("/api/v1/productos/").json()
        self.assertEqual(data["count"], 1)
        self.assertNotIn("reference_price", data["results"][0])
        response = self.client.post("/api/v1/productos/", {"sku": "X"}, content_type="application/json")
        self.assertEqual(response.status_code, 403)

    def test_admin_creates_and_cannot_delete(self):
        self.client.force_login(make_user("adm", ROLE_ADMIN))
        payload = {
            "sku": "api-2",
            "name": "Nuevo",
            "category": self.category.pk,
            "unit": UnitOfMeasure.objects.get(code="PZA").pk,
            "reference_price": "12.50",
        }
        response = self.client.post("/api/v1/productos/", payload, content_type="application/json")
        self.assertEqual(response.status_code, 201, response.content)
        self.assertTrue(Product.objects.filter(sku="API-2").exists())
        self.assertTrue(AuditLog.objects.filter(action="create", object_repr__startswith="API-2").exists())
        self.assertEqual(self.client.delete(f"/api/v1/productos/{self.product.pk}/").status_code, 405)
