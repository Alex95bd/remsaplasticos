from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="0A2540")


def export_stock(queryset) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Existencias"
    headers = ["SKU", "Producto", "Almacén", "Ubicación", "Unidad", "Existencia", "Stock mínimo", "Bajo mínimo"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for s in queryset:
        ws.append(
            [
                s.product.sku,
                s.product.name,
                s.warehouse.code,
                s.location.code if s.location else "",
                s.product.unit.code,
                s.quantity,
                s.min_quantity,
                "Sí" if s.is_low else "",
            ]
        )
    for col, width in zip("ABCDEFGH", [16, 40, 10, 12, 8, 14, 14, 12], strict=True):
        ws.column_dimensions[col].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
